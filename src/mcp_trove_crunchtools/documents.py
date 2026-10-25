"""Readers for office documents, spreadsheets and presentations.

Each takes a path and returns the text to index, raising the underlying
library's error, or ValueError naming the problem, when the file cannot be
read; extract_text turns either into an ExtractionError for that file. Like
PDFs, spreadsheets are indexed by their opening rows: enough to know what
a sheet is about, not every number in it.
"""

from __future__ import annotations

import datetime
import struct
import zipfile
from typing import TYPE_CHECKING, Any

from defusedxml import ElementTree

if TYPE_CHECKING:
    from collections.abc import Callable, Iterable
    from pathlib import Path

SHEET_MAX_ROWS = 200  # leading non-empty rows of each sheet indexed
ODF_MAX_XML_BYTES = 50 * 1024 * 1024  # an ODF content.xml larger than this is refused
# docx, xlsx, pptx and the OpenDocument formats are zip packages. One whose
# members unpack to more than this is refused before it is opened: a small
# file can expand without bound, and the indexer runs in 1.5 GB.
PACKAGE_MAX_UNZIPPED = 256 * 1024 * 1024

_TEXT_NS = "urn:oasis:names:tc:opendocument:xmlns:text:1.0"
_TABLE_NS = "urn:oasis:names:tc:opendocument:xmlns:table:1.0"
_DRAW_NS = "urn:oasis:names:tc:opendocument:xmlns:drawing:1.0"
_PRESENTATION_NS = "urn:oasis:names:tc:opendocument:xmlns:presentation:1.0"
_PARAGRAPHS = {f"{{{_TEXT_NS}}}p", f"{{{_TEXT_NS}}}h"}
# Placeholders a slide inherits from its master: their text is a field
# marker such as "<number>", not content.
_PLACEHOLDER_CLASSES = {"page-number", "date-time", "footer", "header"}


def _cell_text(value: Any) -> str:
    """A cell as text; a date with no time of day loses the midnight."""
    if isinstance(value, datetime.datetime) and value.time() == datetime.time():
        return value.date().isoformat()
    return str(value).strip()


def _rows_text(title: str, rows: Iterable[Iterable[Any]]) -> str:
    """A sheet as a heading and up to SHEET_MAX_ROWS tab-separated, non-empty rows."""
    lines = [f"[Sheet] {title}"]
    for row in rows:
        cells = [text for value in row if value is not None and (text := _cell_text(value))]
        if cells:
            lines.append("\t".join(cells))
            if len(lines) > SHEET_MAX_ROWS:
                break
    return "\n".join(lines)


def _check_package(path: Path) -> None:
    """Refuse a zip package whose members would unpack past PACKAGE_MAX_UNZIPPED."""
    with zipfile.ZipFile(path) as archive:
        unzipped = sum(info.file_size for info in archive.infolist())
    if unzipped > PACKAGE_MAX_UNZIPPED:
        msg = f"unpacks to {unzipped} bytes, over the {PACKAGE_MAX_UNZIPPED} byte limit"
        raise ValueError(msg)


def read_docx(path: Path) -> str:
    import docx

    _check_package(path)
    document = docx.Document(str(path))
    return "\n\n".join(p.text for p in document.paragraphs if p.text.strip())


def read_xlsx(path: Path) -> str:
    import openpyxl

    _check_package(path)
    workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    try:
        return "\n\n".join(
            _rows_text(sheet.title, sheet.iter_rows(values_only=True))
            for sheet in workbook.worksheets
        )
    finally:
        workbook.close()


def read_xls(path: Path) -> str:
    import xlrd

    workbook = xlrd.open_workbook(str(path), on_demand=True)
    try:
        sheets = []
        for index in range(workbook.nsheets):
            sheet = workbook.sheet_by_index(index)
            rows = (
                [_xls_value(cell, workbook.datemode) for cell in sheet.row(r)]
                for r in range(sheet.nrows)
            )
            sheets.append(_rows_text(sheet.name, rows))
            workbook.unload_sheet(index)
        return "\n\n".join(sheets)
    finally:
        workbook.release_resources()


def _xls_value(cell: Any, datemode: int) -> Any:
    """An .xls cell's value; dates are stored as day counts and converted."""
    import xlrd

    if cell.ctype == xlrd.XL_CELL_DATE:
        return xlrd.xldate_as_datetime(cell.value, datemode)
    return cell.value


def read_table(path: Path) -> str:
    """A CSV or TSV file's leading lines."""
    lines: list[str] = []
    with path.open(encoding="utf-8", errors="replace") as handle:
        for line in handle:
            if line.strip():
                lines.append(line.rstrip("\n"))
                if len(lines) >= SHEET_MAX_ROWS:
                    break
    return "\n".join(lines)


def read_pptx(path: Path) -> str:
    """Slide text and speaker notes; images are not looked at."""
    from pptx import Presentation

    _check_package(path)
    slides = []
    for number, slide in enumerate(Presentation(str(path)).slides, start=1):
        texts = [f"[Slide {number}]", *_shape_texts(slide.shapes)]
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame
            if notes is not None and notes.text.strip():
                texts.append(f"[Notes] {notes.text.strip()}")
        slides.append("\n".join(texts))
    return "\n\n".join(slides)


def _shape_texts(shapes: Iterable[Any]) -> list[str]:
    texts: list[str] = []
    for shape in shapes:
        if hasattr(shape, "shapes"):  # a group
            texts.extend(_shape_texts(shape.shapes))
        elif shape.has_text_frame and shape.text_frame.text.strip():
            texts.append(shape.text_frame.text.strip())
        elif getattr(shape, "has_table", False):
            for row in shape.table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    texts.append("\t".join(cells))
    return texts


def read_rtf(path: Path) -> str:
    from striprtf.striprtf import rtf_to_text

    return str(rtf_to_text(path.read_text(encoding="latin-1"), errors="ignore"))


def read_odf(path: Path) -> str:
    """Text of an OpenDocument file: paragraphs, sheets (leading rows) or slides."""
    _check_package(path)
    with zipfile.ZipFile(path) as archive:
        if _odf_encrypted(archive):
            msg = "document is password-protected"
            raise ValueError(msg)
        info = archive.getinfo("content.xml")
        if info.file_size > ODF_MAX_XML_BYTES:
            msg = f"content.xml is {info.file_size} bytes, over {ODF_MAX_XML_BYTES}"
            raise ValueError(msg)
        root = ElementTree.fromstring(archive.read(info))
    tables = list(root.iter(f"{{{_TABLE_NS}}}table"))
    if path.suffix.lower() == ".ods":
        return "\n\n".join(_odf_sheet(table) for table in tables)
    pages = list(root.iter(f"{{{_DRAW_NS}}}page"))
    if pages:
        return "\n\n".join(
            "\n".join([f"[Slide {number}]", *_odf_paragraphs(page)])
            for number, page in enumerate(pages, start=1)
        )
    return "\n\n".join(_odf_paragraphs(root))


def _odf_encrypted(archive: zipfile.ZipFile) -> bool:
    try:
        manifest = archive.read("META-INF/manifest.xml")
    except KeyError:
        return False
    return b"encryption-data" in manifest


def _odf_paragraphs(element: Any) -> list[str]:
    """Paragraph and heading text under element, in order, skipping placeholders."""
    texts: list[str] = []
    for child in element:
        if child.get(f"{{{_PRESENTATION_NS}}}class") in _PLACEHOLDER_CLASSES:
            continue
        if child.tag in _PARAGRAPHS:
            if text := "".join(child.itertext()).strip():
                texts.append(text)
        else:
            texts.extend(_odf_paragraphs(child))
    return texts


def _odf_sheet(table: Any) -> str:
    rows = (
        ["".join(cell.itertext()) for cell in row]
        for row in table.iter(f"{{{_TABLE_NS}}}table-row")
    )
    return _rows_text(table.get(f"{{{_TABLE_NS}}}name", "Sheet"), rows)


# Word 97-2003 (.doc), read from its piece table ([MS-DOC] 2.4.1), so no
# converter binary has to ship in the image.
_WORD_IDENT = 0xA5EC
_WORD97_MIN_NFIB = 0x00C1
_FIB_ENCRYPTED = 0x0100
_FIB_TABLE_1 = 0x0200
_CLX_PRC = 0x01  # a Prc block (property modifiers) ahead of the piece table
_CLX_PCDT = 0x02  # the piece table itself
_PLC_ENTRY = 12  # each piece adds a 4-byte CP and an 8-byte Pcd to the PlcPcd
_FCCLX_INDEX = 33  # fcClx/lcbClx's pair in FibRgFcLcb97
_CCPTEXT_INDEX = 3  # ccpText in FibRgLw97
_FC_COMPRESSED = 0x40000000
_WORD_CONTROL = str.maketrans({"\r": "\n", "\x07": "\t", "\x0b": "\n", "\x0c": "\n", "\x1e": "-"})


def read_doc(path: Path) -> str:
    import olefile

    with olefile.OleFileIO(str(path)) as ole:
        if not ole.exists("WordDocument"):
            msg = "not a Word document (no WordDocument stream)"
            raise ValueError(msg)
        word = ole.openstream("WordDocument").read()
        ident, nfib = struct.unpack_from("<HH", word, 0)
        (flags,) = struct.unpack_from("<H", word, 0x0A)
        if ident != _WORD_IDENT:
            msg = "not a Word document (bad FIB signature)"
            raise ValueError(msg)
        if nfib < _WORD97_MIN_NFIB:
            msg = "Word 95 or older format is not supported"
            raise ValueError(msg)
        if flags & _FIB_ENCRYPTED:
            msg = "document is password-protected"
            raise ValueError(msg)
        table_name = "1Table" if flags & _FIB_TABLE_1 else "0Table"
        table = ole.openstream(table_name).read()
    fc_clx, lcb_clx, ccp_text = _fib_fields(word)
    pieces = _piece_table(table[fc_clx : fc_clx + lcb_clx])
    return _clean_word_text(_piece_text(word, pieces, ccp_text))


def _fib_fields(word: bytes) -> tuple[int, int, int]:
    """fcClx, lcbClx and ccpText from the FIB's variable-length arrays."""
    offset = 32
    (csw,) = struct.unpack_from("<H", word, offset)
    offset += 2 + csw * 2
    (cslw,) = struct.unpack_from("<H", word, offset)
    (ccp_text,) = struct.unpack_from("<i", word, offset + 2 + _CCPTEXT_INDEX * 4)
    offset += 2 + cslw * 4 + 2  # past cbRgFcLcb
    fc_clx, lcb_clx = struct.unpack_from("<II", word, offset + _FCCLX_INDEX * 8)
    return fc_clx, lcb_clx, ccp_text


def _piece_table(clx: bytes) -> list[tuple[int, int, int]]:
    """(cp_start, cp_end, fc) for each piece, skipping any leading Prc blocks."""
    offset = 0
    while offset < len(clx) and clx[offset] == _CLX_PRC:
        (cb_grpprl,) = struct.unpack_from("<H", clx, offset + 1)
        offset += 3 + cb_grpprl
    if offset >= len(clx) or clx[offset] != _CLX_PCDT:
        msg = "piece table not found"
        raise ValueError(msg)
    (lcb,) = struct.unpack_from("<I", clx, offset + 1)
    plc = clx[offset + 5 : offset + 5 + lcb]
    count = (lcb - 4) // _PLC_ENTRY
    cps = struct.unpack_from(f"<{count + 1}i", plc, 0)
    pieces = []
    for index in range(count):
        (fc,) = struct.unpack_from("<I", plc, (count + 1) * 4 + index * 8 + 2)
        pieces.append((cps[index], cps[index + 1], fc))
    return pieces


def _piece_text(word: bytes, pieces: list[tuple[int, int, int]], ccp_text: int) -> str:
    """The main document's characters, in order, up to ccpText."""
    parts = []
    for cp_start, cp_end, fc in pieces:
        if cp_start >= ccp_text:
            break
        length = min(cp_end, ccp_text) - cp_start
        if fc & _FC_COMPRESSED:
            start = (fc & ~_FC_COMPRESSED) // 2
            parts.append(word[start : start + length].decode("cp1252", errors="replace"))
        else:
            parts.append(word[fc : fc + 2 * length].decode("utf-16-le", errors="replace"))
    return "".join(parts)


def _clean_word_text(text: str) -> str:
    """Drop field instructions (keeping their results) and control characters."""
    kept: list[str] = []
    # One entry per open field: True while in its instruction part.
    fields: list[bool] = []
    for char in text.translate(_WORD_CONTROL):
        match char:
            case "\x13":  # field start: its instruction follows
                fields.append(True)
            case "\x14" if fields:  # separator: the field's result follows
                fields[-1] = False
            case "\x15" if fields:  # field end
                fields.pop()
            case _ if not any(fields) and (char >= " " or char in "\n\t"):
                kept.append(char)
    lines = (line.rstrip() for line in "".join(kept).split("\n"))
    return "\n".join(line for line in lines if line)


DOCUMENT_EXTRACTORS: dict[str, Callable[[Path], str]] = {
    "docx": read_docx,
    "doc": read_doc,
    "odf": read_odf,
    "rtf": read_rtf,
    "xlsx": read_xlsx,
    "xls": read_xls,
    "table": read_table,
    "pptx": read_pptx,
}
